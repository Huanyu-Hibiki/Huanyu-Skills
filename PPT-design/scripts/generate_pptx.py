"""Render a confirmed PPTDESIGN.md into an editable PowerPoint deck."""

from __future__ import annotations

from dataclasses import dataclass
import argparse
from html import escape
import math
import os
from pathlib import Path
import re
import tempfile

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.util import Inches, Pt

from common import atomic_write_text, clean_text, safe_output_dir


@dataclass(frozen=True)
class SlideSpec:
    number: int
    title: str
    takeaway: str
    content: str
    editable_objects: str
    table_data: tuple[tuple[str, ...], ...]
    chart_data: tuple[tuple[str, float], ...]


def generate_pptx(
    design_path: Path,
    output_path: Path,
    *,
    allowed_root: Path | None = None,
) -> Path:
    """Generate an object-level PPTX from a confirmed design document."""

    design_path = _validate_file_path(design_path, ".md")
    output_path = _validate_output_path(output_path, allowed_root)
    text = design_path.read_text(encoding="utf-8")
    metadata = _read_frontmatter(text)
    if metadata.get("status") != "confirmed":
        raise ValueError("PPTDESIGN.md must have status: confirmed before generating")
    if metadata.get("editable_output", "").lower() not in {"true", "yes"}:
        raise ValueError("PPTDESIGN.md must allow editable output")

    slides = _read_slides(text)
    if not slides:
        raise ValueError("PPTDESIGN.md does not contain a slide outline")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]
    for spec in slides:
        slide = presentation.slides.add_slide(blank_layout)
        _render_slide(slide, spec)

    preview_path = output_path.with_suffix(".preview.html")
    if output_path.is_symlink():
        raise ValueError("output PPTX must not be a symbolic link")
    if preview_path.is_symlink():
        raise ValueError("preview file must not be a symbolic link")
    _reject_linked_file(output_path, "output PPTX")
    _reject_linked_file(preview_path, "preview file")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    _atomic_save_presentation(presentation, output_path)
    _write_preview(preview_path, slides)
    return output_path


def _validate_file_path(path: Path, suffix: str) -> Path:
    if not isinstance(path, Path):
        raise TypeError("path must be a pathlib.Path")
    resolved = path.absolute()
    if resolved.suffix.lower() != suffix:
        raise ValueError(f"design file must use {suffix}")
    if not resolved.is_file() or resolved.is_symlink():
        raise ValueError("design file must be a regular file")
    return resolved


def _validate_output_path(path: Path, allowed_root: Path | None) -> Path:
    if not isinstance(path, Path):
        raise TypeError("output_path must be a pathlib.Path")
    if path.suffix.lower() != ".pptx":
        raise ValueError("output file must use .pptx")
    if ".." in path.parts:
        raise ValueError("output_path must not contain parent traversal")
    parent = safe_output_dir(path.parent, allowed_root=allowed_root)
    return parent / path.name


def _reject_linked_file(path: Path, label: str) -> None:
    """Avoid overwriting an existing hard-linked artifact."""

    if path.exists() and path.is_file() and path.stat().st_nlink > 1:
        raise ValueError(f"{label} must not be a hard link")


def _atomic_save_presentation(presentation: PresentationType, output_path: Path) -> None:
    temporary: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            suffix=".pptx",
            dir=output_path.parent,
            prefix=f".{output_path.stem}.",
            delete=False,
        ) as handle:
            temporary = Path(handle.name)
        presentation.save(str(temporary))
        os.replace(temporary, output_path)
        temporary = None
    finally:
        if temporary is not None:
            try:
                temporary.unlink()
            except FileNotFoundError:
                pass


def _read_frontmatter(text: str) -> dict[str, str]:
    if not text.startswith("---\n"):
        raise ValueError("PPTDESIGN.md must start with YAML frontmatter")
    match = re.match(r"\A---\r?\n(.*?)\r?\n---\r?\n", text, re.DOTALL)
    if not match:
        raise ValueError("PPTDESIGN.md has incomplete frontmatter")
    values: dict[str, str] = {}
    for line in match.group(1).splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        values[key.strip()] = value.strip().strip('"\'')
    return values


def _read_slides(text: str) -> list[SlideSpec]:
    sections = list(re.finditer(r"(?m)^### 第 (\d+) 页\s*$", text))
    slides: list[SlideSpec] = []
    for index, match in enumerate(sections):
        end = sections[index + 1].start() if index + 1 < len(sections) else len(text)
        block = text[match.end() : end]
        fields = _read_slide_fields(block)
        slides.append(
            SlideSpec(
                number=int(match.group(1)),
                title=fields.get("标题", f"第 {match.group(1)} 页"),
                takeaway=fields.get("核心结论", ""),
                content=fields.get("内容", ""),
                editable_objects=fields.get("可编辑对象", ""),
                table_data=_parse_table_data(fields.get("表格数据", "")),
                chart_data=_parse_chart_data(fields.get("图表数据", "")),
            )
        )
    return slides


def _read_slide_fields(block: str) -> dict[str, str]:
    fields: dict[str, str] = {}
    for line in block.splitlines():
        match = re.match(r"^-\s*([^：:]+)[：:]\s*(.*)$", line.strip())
        if match:
            fields[match.group(1).strip()] = match.group(2).strip()
    return fields


def _parse_table_data(raw: str) -> tuple[tuple[str, ...], ...]:
    if not raw:
        return ()
    rows = []
    for row in raw.split(";"):
        cells = tuple(clean_text(cell, "table cell", 20) for cell in row.split(","))
        if len(cells) < 2:
            raise ValueError("table data rows need at least two cells")
        rows.append(cells)
    width = len(rows[0])
    if len(rows) < 2 or any(len(row) != width for row in rows):
        raise ValueError("table data must be a rectangular matrix")
    if len(rows) > 4 or width > 6:
        raise ValueError("table data exceeds the readable 4x6 layout limit")
    return tuple(rows)


def _parse_chart_data(raw: str) -> tuple[tuple[str, float], ...]:
    if not raw:
        return ()
    values = []
    for item in raw.split(","):
        if "=" not in item:
            raise ValueError("chart data must use label=value pairs")
        label, value = item.split("=", 1)
        number = float(value.strip())
        if not math.isfinite(number):
            raise ValueError("chart data values must be finite")
        values.append((clean_text(label, "chart label", 100), number))
    if len(values) < 2:
        raise ValueError("chart data needs at least two values")
    if len(values) > 12:
        raise ValueError("chart data exceeds the readable 12-point layout limit")
    return tuple(values)


def _render_slide(slide: Slide, spec: SlideSpec) -> None:
    # python-pptx's slide object is intentionally kept behind this small
    # renderer seam; every visible item below is a native PowerPoint shape.
    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.55), Inches(11.9), Inches(0.7))
    title = title_box.text_frame.paragraphs[0]
    title.text = clean_text(spec.title, "slide title", 120)
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(20, 31, 48)

    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.4), Inches(1.0), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(30, 110, 180)
    accent.line.fill.background()

    objects = spec.editable_objects
    has_data_visual = "表格" in objects or "图表" in objects
    body_height = 2.8 if has_data_visual else 3.9
    takeaway = clean_text(spec.takeaway or spec.content, "slide content", 120 if has_data_visual else 180)
    detail_text = clean_text(spec.content, "slide detail", 120 if has_data_visual else 180) if spec.takeaway and spec.content else ""
    if len(takeaway) + len(detail_text) > (200 if has_data_visual else 300):
        raise ValueError(f"slide {spec.number} text exceeds the readable layout limit")
    body_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.85), Inches(11.4), Inches(body_height))
    frame = body_box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = takeaway
    paragraph.font.size = Pt(22)
    paragraph.font.color.rgb = RGBColor(45, 55, 72)
    paragraph.alignment = PP_ALIGN.LEFT
    if detail_text:
        detail = frame.add_paragraph()
        detail.text = detail_text
        detail.font.size = Pt(16)
        detail.font.color.rgb = RGBColor(80, 90, 105)

    if "表格" in objects:
        if not spec.table_data:
            raise ValueError(f"slide {spec.number} requests a table without 表格数据")
        table_data = spec.table_data
        table_height = max(0.85, min(1.2, 0.3 * len(table_data)))
        table = slide.shapes.add_table(
            len(table_data), len(table_data[0]), Inches(0.9), Inches(5.25), Inches(5.0), Inches(table_height)
        ).table
        for row in table.rows:
            row.height = Inches(table_height / len(table_data))
        for row, cells in enumerate(table_data):
            for column, value in enumerate(cells):
                cell = table.cell(row, column)
                cell.text = value
                cell.text_frame.word_wrap = True
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(10)
    if "图表" in objects:
        if not spec.chart_data:
            raise ValueError(f"slide {spec.number} requests a chart without 图表数据")
        chart_data = ChartData()
        chart_data.categories = [label for label, _ in spec.chart_data]
        chart_data.add_series("指标", [value for _, value in spec.chart_data])
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(6.3),
            Inches(4.85),
            Inches(5.2),
            Inches(1.5),
            chart_data,
        )

    footer = slide.shapes.add_textbox(Inches(11.7), Inches(6.85), Inches(0.9), Inches(0.3))
    footer.text_frame.paragraphs[0].text = str(spec.number)
    footer.text_frame.paragraphs[0].font.size = Pt(10)
    footer.text_frame.paragraphs[0].font.color.rgb = RGBColor(100, 110, 125)


def _write_preview(path: Path, slides: list[SlideSpec]) -> None:
    cards = "\n".join(
        "<article class='slide'>"
        f"<div class='page'>{spec.number}</div>"
        f"<h1>{escape(spec.title)}</h1>"
        f"<p>{escape(spec.takeaway)}</p>"
        f"<p class='detail'>{escape(spec.content)}</p>"
        f"<p class='object-plan'>可编辑对象：{escape(spec.editable_objects or '文本框和形状')}</p>"
        f"{_preview_objects(spec)}"
        "</article>"
        for spec in slides
    )
    html = f"""<!doctype html>
<meta charset="utf-8">
<title>PPT preview</title>
<style>
body {{ margin: 0; padding: 24px; background: #e9edf2; font-family: Arial, sans-serif; }}
.slide {{ width: 960px; min-height: 540px; height: auto; margin: 0 auto 24px; padding: 48px; box-sizing: border-box; background: white; color: #141f30; position: relative; overflow: visible; overflow-wrap: anywhere; }}
h1 {{ font-size: 40px; margin: 0 0 56px; }} p {{ font-size: 26px; color: #2d3748; }} .page {{ position: absolute; right: 48px; bottom: 24px; color: #647080; }}
table.preview-table {{ border-collapse: collapse; font-size: 14px; }} table.preview-table td {{ border: 1px solid #b8c2cf; padding: 4px 8px; }}
.chart-data {{ display: flex; gap: 8px; align-items: center; height: 70px; margin-top: 8px; border-bottom: 1px solid #9aa8b8; }} .bar {{ min-width: 34px; background: #1e6eb4; color: white; font-size: 11px; text-align: center; }} .bar.negative {{ background: #b94b5b; }}
</style>
{cards}
"""
    atomic_write_text(path, html)


def _preview_objects(spec: SlideSpec) -> str:
    chunks: list[str] = []
    if spec.table_data:
        rows = "".join(
            "<tr>" + "".join(f"<td>{escape(cell)}</td>" for cell in row) + "</tr>"
            for row in spec.table_data
        )
        chunks.append(f"<table class='preview-table'>{rows}</table>")
    if spec.chart_data:
        maximum = max(abs(value) for _, value in spec.chart_data) or 1.0
        bars = "".join(
            f"<div class='bar{' negative' if value < 0 else ''}' style='height:{max(8, int(abs(value) / maximum * 60))}px' title='{escape(label)}'>{escape(str(value))}</div>"
            for label, value in spec.chart_data
        )
        chunks.append(f"<div class='chart-data'>{bars}</div>")
    return "".join(chunks)


def main() -> int:
    parser = argparse.ArgumentParser(description="Render confirmed PPTDESIGN.md to editable PPTX")
    parser.add_argument("design", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    generate_pptx(args.design, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
