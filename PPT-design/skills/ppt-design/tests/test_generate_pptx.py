"""End-to-end tests for the native PPTX generation seam."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import sys

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from generate_pptx import generate_pptx  # noqa: E402


DESIGN = """---
name: "Editable demo"
mode: topic-interview
status: confirmed
editable_output: true
---

# PPTDESIGN: Editable demo

## 叙事
- 主题：Editable demo

## 视觉 tokens
- 画幅：16:9

## 来源
- 用户材料：演示数据

## 逐页大纲
### 第 1 页
- 标题：结论
- 核心结论：文字保持可编辑
- 内容：第一张演示页
- 可编辑对象：表格和图表
- 表格数据：指标,数值;收入,20
- 图表数据：当前=1,目标=2
### 第 2 页
- 标题：证据
- 核心结论：形状保持可编辑
- 内容：第二张演示页

## 待确认项
- 已确认
"""


def test_unconfirmed_design_is_rejected(tmp_path: Path) -> None:
    design = tmp_path / "PPTDESIGN.md"
    design.write_text(DESIGN.replace("status: confirmed", "status: awaiting-user-confirmation"), encoding="utf-8")

    with pytest.raises(ValueError, match="confirmed"):
        generate_pptx(design, tmp_path / "deck.pptx")


def test_missing_editable_output_metadata_is_rejected(tmp_path: Path) -> None:
    design = tmp_path / "PPTDESIGN.md"
    design.write_text(DESIGN.replace("editable_output: true\n", ""), encoding="utf-8")

    with pytest.raises(ValueError, match="editable output"):
        generate_pptx(design, tmp_path / "deck.pptx")


def test_confirmed_design_generates_native_editable_pptx(tmp_path: Path) -> None:
    design = tmp_path / "PPTDESIGN.md"
    output = tmp_path / "deck.pptx"
    design.write_text(DESIGN, encoding="utf-8")

    result = generate_pptx(design, output)

    assert result == output
    assert (tmp_path / "deck.preview.html").is_file()
    preview = (tmp_path / "deck.preview.html").read_text(encoding="utf-8")
    assert "文字保持可编辑" in preview
    assert "表格和图表" in preview
    assert "收入" in preview
    assert "2.0" in preview
    presentation = Presentation(str(output))
    assert len(presentation.slides) == 2
    assert any(
        shape.has_text_frame
        for slide in presentation.slides
        for shape in slide.shapes
    )
    assert any(
        shape.shape_type is not MSO_SHAPE_TYPE.PICTURE
        for slide in presentation.slides
        for shape in slide.shapes
    )
    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    )
    assert "结论" in text
    assert "文字保持可编辑" in text
    assert "第一张演示页" in text
    assert "第二张演示页" in text
    assert any(shape.has_table for slide in presentation.slides for shape in slide.shapes)
    assert any(shape.has_chart for slide in presentation.slides for shape in slide.shapes)
    assert not any(
        shape.shape_type is MSO_SHAPE_TYPE.PICTURE
        for slide in presentation.slides
        for shape in slide.shapes
    )
    table_shapes = [
        shape for slide in presentation.slides for shape in slide.shapes if shape.has_table
    ]
    assert table_shapes[0].table.cell(1, 0).text == "收入"
    assert table_shapes[0].table.cell(1, 1).text == "20"


def test_text_that_cannot_fit_the_fixed_layout_is_rejected(tmp_path: Path) -> None:
    design = tmp_path / "PPTDESIGN.md"
    long_content = "内容" * 500
    design.write_text(DESIGN.replace("第二张演示页", long_content), encoding="utf-8")

    with pytest.raises(ValueError, match="character limit|layout limit"):
        generate_pptx(design, tmp_path / "deck.pptx")


def test_preview_preserves_negative_chart_values(tmp_path: Path) -> None:
    design = tmp_path / "PPTDESIGN.md"
    design.write_text(DESIGN.replace("当前=1,目标=2", "下降=-2,目标=2"), encoding="utf-8")

    generate_pptx(design, tmp_path / "deck.pptx")

    preview = (tmp_path / "deck.preview.html").read_text(encoding="utf-8")
    assert "negative" in preview
    assert "-2.0" in preview
